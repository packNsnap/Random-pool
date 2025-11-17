from fastapi import FastAPI, Request, Depends, Form, UploadFile, File, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from starlette.middleware.sessions import SessionMiddleware
from sqlalchemy.orm import Session
from datetime import datetime
import os
import shutil
import uuid
import re
import csv
import io
from typing import Optional

from database import get_db, init_db
from models import User, Client, Roster, EmailTemplate, EmailLog, Settings, Attachment, RosterEntry, CCEmail
from auth import authenticate_user, require_login, get_current_user, hash_password
from email_service import EmailService
from scheduler import start_scheduler
from utils import get_current_quarter, format_datetime, days_since, render_template_string

app = FastAPI(title="Random Pool Management")

SESSION_SECRET = os.getenv("SESSION_SECRET")
if not SESSION_SECRET:
    raise RuntimeError("SESSION_SECRET environment variable must be set")

app.add_middleware(SessionMiddleware, secret_key=SESSION_SECRET)

templates = Jinja2Templates(directory="templates")
templates.env.filters['format_datetime'] = format_datetime
templates.env.filters['days_since'] = days_since

os.makedirs("static", exist_ok=True)
os.makedirs("uploads/rosters", exist_ok=True)
os.makedirs("uploads/attachments", exist_ok=True)

try:
    app.mount("/static", StaticFiles(directory="static"), name="static")
except Exception:
    pass

@app.on_event("startup")
async def startup_event():
    init_db()
    start_scheduler()

@app.get("/", response_class=HTMLResponse)
async def dashboard(request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    if not user:
        return RedirectResponse(url="/login", status_code=303)
    
    clients = db.query(Client).filter(Client.active == True).all()
    current_quarter = get_current_quarter()
    
    client_data = []
    for client in clients:
        roster_received = any(r.quarter == current_quarter for r in client.rosters)
        days_last_contact = days_since(client.last_roster_request_at)
        
        client_data.append({
            "client": client,
            "roster_received": roster_received,
            "days_last_contact": days_last_contact
        })
    
    return templates.TemplateResponse("dashboard.html", {
        "request": request,
        "user": user,
        "client_data": client_data,
        "current_quarter": current_quarter
    })

@app.get("/login", response_class=HTMLResponse)
async def login_page(request: Request):
    return templates.TemplateResponse("login.html", {"request": request, "error": None})

@app.post("/login")
async def login(request: Request, username: str = Form(...), password: str = Form(...), db: Session = Depends(get_db)):
    user = authenticate_user(db, username, password)
    if not user:
        return templates.TemplateResponse("login.html", {
            "request": request,
            "error": "Invalid username or password"
        })
    
    request.session["user_id"] = user.id
    return RedirectResponse(url="/", status_code=303)

@app.get("/logout")
async def logout(request: Request):
    request.session.clear()
    return RedirectResponse(url="/login", status_code=303)

@app.get("/clients", response_class=HTMLResponse)
async def list_clients(request: Request, user: User = Depends(require_login), db: Session = Depends(get_db)):
    clients = db.query(Client).all()
    return templates.TemplateResponse("clients.html", {
        "request": request,
        "user": user,
        "clients": clients
    })

@app.get("/clients/new", response_class=HTMLResponse)
async def new_client_form(request: Request, user: User = Depends(require_login)):
    return templates.TemplateResponse("client_form.html", {
        "request": request,
        "user": user,
        "client": None
    })

@app.post("/clients/new")
async def create_client(
    request: Request,
    name: str = Form(...),
    contact_name: str = Form(...),
    contact_email: str = Form(...),
    program_type: str = Form("DOT"),
    progress_frequency_days: int = Form(14),
    reminder_enabled: Optional[str] = Form(None),
    reminder_frequency_days: int = Form(7),
    testing_report_enabled: Optional[str] = Form(None),
    testing_report_day_of_week: int = Form(1),
    active: Optional[str] = Form(None),
    notes: Optional[str] = Form(None),
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    client = Client(
        name=name,
        contact_name=contact_name,
        contact_email=contact_email,
        program_type=program_type,
        progress_frequency_days=progress_frequency_days,
        reminder_enabled=reminder_enabled == "true",
        reminder_frequency_days=reminder_frequency_days,
        testing_report_enabled=testing_report_enabled == "true",
        testing_report_day_of_week=testing_report_day_of_week,
        active=active == "true",
        notes=notes
    )
    db.add(client)
    db.commit()
    return RedirectResponse(url="/clients", status_code=303)

@app.get("/clients/{client_id}", response_class=HTMLResponse)
async def client_detail(
    request: Request,
    client_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    rosters = sorted(client.rosters, key=lambda r: r.received_at, reverse=True)
    email_logs = db.query(EmailLog).filter(EmailLog.client_id == client_id).order_by(EmailLog.sent_at.desc()).limit(10).all()
    active_templates = db.query(EmailTemplate).filter(EmailTemplate.active == True).all()
    
    return templates.TemplateResponse("client_detail.html", {
        "request": request,
        "user": user,
        "client": client,
        "rosters": rosters,
        "email_logs": email_logs,
        "current_quarter": get_current_quarter(),
        "active_templates": active_templates
    })

@app.get("/clients/{client_id}/edit", response_class=HTMLResponse)
async def edit_client_form(
    request: Request,
    client_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    return templates.TemplateResponse("client_form.html", {
        "request": request,
        "user": user,
        "client": client
    })

@app.post("/clients/{client_id}/edit")
async def update_client(
    request: Request,
    client_id: int,
    name: str = Form(...),
    contact_name: str = Form(...),
    contact_email: str = Form(...),
    program_type: str = Form("DOT"),
    progress_frequency_days: int = Form(14),
    reminder_enabled: Optional[str] = Form(None),
    reminder_frequency_days: int = Form(7),
    testing_report_enabled: Optional[str] = Form(None),
    testing_report_day_of_week: int = Form(1),
    active: Optional[str] = Form(None),
    notes: Optional[str] = Form(None),
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    client.name = name
    client.contact_name = contact_name
    client.contact_email = contact_email
    client.program_type = program_type
    client.progress_frequency_days = progress_frequency_days
    client.reminder_enabled = reminder_enabled == "true"
    client.reminder_frequency_days = reminder_frequency_days
    client.testing_report_enabled = testing_report_enabled == "true"
    client.testing_report_day_of_week = testing_report_day_of_week
    client.active = active == "true"
    client.notes = notes
    db.commit()
    
    return RedirectResponse(url=f"/clients/{client_id}", status_code=303)

@app.post("/clients/{client_id}/send_roster_request")
async def send_roster_request(
    client_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    template = db.query(EmailTemplate).filter(
        EmailTemplate.template_type == "roster_request",
        EmailTemplate.active == True
    ).first()
    
    if not template:
        raise HTTPException(status_code=404, detail="No roster request template found")
    
    current_quarter = get_current_quarter()
    variables = {
        "client_name": client.name,
        "contact_name": client.contact_name,
        "quarter": current_quarter,
        "due_date": datetime.utcnow().strftime("%B %d, %Y")
    }
    
    attachment_paths = [os.path.join("uploads", "attachments", att.filename) 
                        for att in client.attachments if att.category == "roster_request"]
    
    email_service = EmailService(db)
    success, message = email_service.send_from_template(client, template, variables, attachments=attachment_paths)
    
    if success:
        client.last_roster_request_at = datetime.utcnow()
        client.status = "Roster Request Sent"
        db.commit()
    
    return RedirectResponse(url=f"/clients/{client_id}", status_code=303)

@app.post("/clients/{client_id}/send_reminder")
async def send_reminder(
    client_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    template = db.query(EmailTemplate).filter(
        EmailTemplate.template_type == "follow_up",
        EmailTemplate.active == True
    ).first()
    
    if not template:
        raise HTTPException(status_code=404, detail="No reminder template found")
    
    current_quarter = get_current_quarter()
    days_overdue = days_since(client.last_roster_request_at) if client.last_roster_request_at else 0
    
    variables = {
        "client_name": client.name,
        "contact_name": client.contact_name,
        "quarter": current_quarter,
        "days_overdue": days_overdue
    }
    
    attachment_paths = [os.path.join("uploads", "attachments", att.filename) 
                        for att in client.attachments if att.category == "reminder"]
    
    email_service = EmailService(db)
    success, message = email_service.send_from_template(client, template, variables, attachments=attachment_paths)
    
    if success:
        client.status = "Reminder Sent"
        db.commit()
    
    return RedirectResponse(url=f"/clients/{client_id}", status_code=303)

@app.post("/clients/{client_id}/send_update")
async def send_update(
    client_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    template = db.query(EmailTemplate).filter(
        EmailTemplate.template_type == "progress_update",
        EmailTemplate.active == True
    ).first()
    
    if not template:
        raise HTTPException(status_code=404, detail="No update template found")
    
    variables = {
        "client_name": client.name,
        "contact_name": client.contact_name,
        "update_date": datetime.utcnow().strftime("%B %d, %Y")
    }
    
    attachment_paths = [os.path.join("uploads", "attachments", att.filename) 
                        for att in client.attachments if att.category == "update"]
    
    email_service = EmailService(db)
    success, message = email_service.send_from_template(client, template, variables, attachments=attachment_paths)
    
    if success:
        client.last_progress_update_at = datetime.utcnow()
        client.status = "Update Sent"
        db.commit()
    
    return RedirectResponse(url=f"/clients/{client_id}", status_code=303)

@app.post("/clients/{client_id}/send_email/{template_id}")
async def send_template_email(
    client_id: int,
    template_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    template = db.query(EmailTemplate).filter(EmailTemplate.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")
    
    roster = db.query(Roster).filter(Roster.client_id == client_id).order_by(Roster.received_at.desc()).first()
    
    variables = {
        "client_name": client.name,
        "contact_name": client.contact_name,
        "quarter": get_current_quarter(),
        "due_date": format_datetime(datetime.utcnow()),
        "days_overdue": days_since(client.last_roster_request_at) if client.last_roster_request_at else 0,
        "update_date": format_datetime(datetime.utcnow()),
        "send_date": format_datetime(datetime.utcnow())
    }
    
    if roster:
        tested_entries = [e for e in roster.entries if e.testing_status == "tested"]
        not_tested_entries = [e for e in roster.entries if e.testing_status == "not_tested"]
        excused_entries = [e for e in roster.entries if e.testing_status == "excused"]
        
        tested_count = len(tested_entries)
        not_tested_count = len(not_tested_entries)
        excused_count = len(excused_entries)
        
        tested_list = "\n".join([f"- {e.employee_name} ({e.employee_id or 'No ID'})" for e in tested_entries]) if tested_entries else "None"
        not_tested_list = "\n".join([f"- {e.employee_name} ({e.employee_id or 'No ID'})" for e in not_tested_entries]) if not_tested_entries else "None"
        excused_list = "\n".join([f"- {e.employee_name} ({e.employee_id or 'No ID'})" for e in excused_entries]) if excused_entries else "None"
        
        tested_percentage = round((tested_count / len(roster.entries) * 100), 1) if len(roster.entries) > 0 else 0
        
        variables.update({
            "roster_quarter": roster.quarter,
            "total_employees": len(roster.entries),
            "tested_count": tested_count,
            "not_tested_count": not_tested_count,
            "excused_count": excused_count,
            "tested_percentage": tested_percentage,
            "tested_list": tested_list,
            "not_tested_list": not_tested_list,
            "excused_list": excused_list
        })
    
    attachment_paths = []
    if template.template_type:
        attachments = db.query(Attachment).filter(
            Attachment.client_id == client_id,
            Attachment.category == template.template_type
        ).all()
        attachment_paths = [os.path.join("uploads", "attachments", a.filename) for a in attachments]
    
    if template.template_type == "progress_update" and roster:
        csv_filename = f"{client.name.replace(' ', '_')}_roster_status_{datetime.utcnow().strftime('%Y%m%d')}.csv"
        csv_path = os.path.join("uploads", "temp", csv_filename)
        os.makedirs("uploads/temp", exist_ok=True)
        
        with open(csv_path, 'w', newline='', encoding='utf-8') as csvfile:
            writer = csv.writer(csvfile)
            writer.writerow(['Employee Name', 'Employee ID', 'Position', 'Department', 'Testing Status', 'Test Date'])
            
            for entry in roster.entries:
                writer.writerow([
                    entry.employee_name,
                    entry.employee_id or '',
                    entry.position or '',
                    entry.department or '',
                    entry.testing_status,
                    entry.test_date.strftime('%Y-%m-%d %H:%M') if entry.test_date else ''
                ])
        
        attachment_paths.append(csv_path)
    
    email_service = EmailService(db)
    success, message = email_service.send_from_template(client, template, variables, attachment_paths)
    
    if template.template_type == "progress_update" and roster:
        if os.path.exists(csv_path):
            os.remove(csv_path)
    
    if success:
        if template.template_type == "roster_request":
            client.last_roster_request_at = datetime.utcnow()
        elif template.template_type == "progress_update":
            client.last_progress_update_at = datetime.utcnow()
        db.commit()
    
    return RedirectResponse(url=f"/clients/{client_id}", status_code=303)

@app.post("/clients/{client_id}/send_quarterly_selections")
async def send_quarterly_selections(
    client_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    template = db.query(EmailTemplate).filter(
        EmailTemplate.template_type == "quarterly_selections",
        EmailTemplate.active == True
    ).first()
    
    if not template:
        raise HTTPException(status_code=404, detail="No quarterly selections template found")
    
    current_quarter = get_current_quarter()
    
    latest_roster = db.query(Roster).filter(
        Roster.client_id == client_id
    ).order_by(Roster.received_at.desc()).first()
    
    roster_quarter = latest_roster.quarter if latest_roster else current_quarter
    
    entries = []
    tested_count = 0
    not_tested_count = 0
    excused_count = 0
    
    if latest_roster:
        entries = db.query(RosterEntry).filter(
            RosterEntry.roster_id == latest_roster.id
        ).all()
        tested_count = sum(1 for e in entries if e.testing_status == "tested")
        excused_count = sum(1 for e in entries if e.testing_status == "excused")
        not_tested_count = sum(1 for e in entries if e.testing_status == "not_tested")
    
    variables = {
        "client_name": client.name,
        "contact_name": client.contact_name,
        "quarter": current_quarter,
        "roster_quarter": roster_quarter,
        "total_employees": len(entries),
        "tested_count": tested_count,
        "not_tested_count": not_tested_count,
        "excused_count": excused_count,
        "send_date": datetime.utcnow().strftime("%B %d, %Y")
    }
    
    attachment_paths = [os.path.join("uploads", "attachments", att.filename) 
                        for att in client.attachments if att.category == "quarterly_selections"]
    
    email_service = EmailService(db)
    success, message = email_service.send_from_template(client, template, variables, attachments=attachment_paths)
    
    if success:
        client.status = "Quarterly Selections Sent"
        db.commit()
    
    return RedirectResponse(url=f"/clients/{client_id}", status_code=303)

@app.post("/clients/{client_id}/mark_roster_received")
async def mark_roster_received(
    client_id: int,
    quarter: str = Form(...),
    file: Optional[UploadFile] = File(None),
    notes: Optional[str] = Form(None),
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    file_path = None
    if file and file.filename:
        file_path = f"uploads/rosters/{client_id}_{quarter}_{file.filename}"
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
    
    roster = Roster(
        client_id=client_id,
        quarter=quarter,
        file_path=file_path,
        notes=notes
    )
    db.add(roster)
    
    client.last_roster_received_at = datetime.utcnow()
    client.status = "Roster Received"
    db.commit()
    
    return RedirectResponse(url=f"/clients/{client_id}", status_code=303)

@app.post("/clients/{client_id}/upload_roster_csv")
async def upload_roster_csv(
    client_id: int,
    quarter: str = Form(...),
    file: UploadFile = File(...),
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    from openpyxl import load_workbook
    
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    if not file.filename:
        raise HTTPException(status_code=400, detail="Please upload a file")
    
    is_csv = file.filename.endswith('.csv')
    is_xlsx = file.filename.endswith('.xlsx') or file.filename.endswith('.xls')
    
    if not (is_csv or is_xlsx):
        raise HTTPException(status_code=400, detail="Please upload a CSV or XLSX file")
    
    content = await file.read()
    file_path = f"uploads/rosters/{client_id}_{quarter}_{file.filename}"
    with open(file_path, "wb") as buffer:
        buffer.write(content)
    
    # Parse the file based on type
    rows = []
    if is_csv:
        csv_data = content.decode('utf-8')
        csv_reader = csv.DictReader(io.StringIO(csv_data))
        rows = list(csv_reader)
    else:  # XLSX
        workbook = load_workbook(file_path)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_dict = {headers[i]: row[i] for i in range(len(headers)) if i < len(row)}
            rows.append(row_dict)
    
    roster = db.query(Roster).filter(
        Roster.client_id == client_id,
        Roster.quarter == quarter
    ).first()
    
    if not roster:
        roster = Roster(
            client_id=client_id,
            quarter=quarter,
            file_path=file_path
        )
        db.add(roster)
        db.flush()
    
    entries_added = 0
    for row in rows:
        normalized_row = {k.lower().strip().replace(' ', '_'): v for k, v in row.items() if k and v}
        
        first_name = normalized_row.get('first_name') or ''
        last_name = normalized_row.get('last_name') or ''
        
        employee_name = (normalized_row.get('employee_name') or 
                        normalized_row.get('name') or 
                        normalized_row.get('full_name') or 
                        normalized_row.get('employee') or 
                        f"{first_name} {last_name}".strip())
        
        if not employee_name:
            continue
        
        entry = RosterEntry(
            roster_id=roster.id,
            client_id=client_id,
            primary_id=normalized_row.get('primary_id'),
            last_name=last_name or None,
            first_name=first_name or None,
            employee_name=employee_name,
            employee_id=(normalized_row.get('employee_id') or 
                        normalized_row.get('id') or 
                        normalized_row.get('emp_id')),
            company=normalized_row.get('company'),
            position=(normalized_row.get('position') or 
                     normalized_row.get('job_title') or 
                     normalized_row.get('title')),
            department=(normalized_row.get('department') or 
                       normalized_row.get('dept')),
            division=normalized_row.get('division'),
            modality=normalized_row.get('modality'),
            location=normalized_row.get('location'),
            supervisor_name=normalized_row.get('supervisor_name'),
            alternate_id=normalized_row.get('alternate_id'),
            alternate_id_type=normalized_row.get('alternate_id_type'),
            alternate_id_2=normalized_row.get('alternate_id_2'),
            alternate_id_2_type=normalized_row.get('alternate_id_2_type'),
            testing_status="not_tested"
        )
        db.add(entry)
        entries_added += 1
    
    client.last_roster_received_at = datetime.utcnow()
    client.status = f"Roster Uploaded ({entries_added} entries)"
    db.commit()
    
    return RedirectResponse(url=f"/clients/{client_id}/roster?roster_id={roster.id}", status_code=303)

@app.post("/clients/{client_id}/toggle_roster_received")
async def toggle_roster_received(
    client_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    if client.last_roster_received_at:
        client.last_roster_received_at = None
        client.status = "Active"
    else:
        client.last_roster_received_at = datetime.utcnow()
        client.status = "Roster Received"
    
    db.commit()
    return RedirectResponse(url=f"/clients/{client_id}", status_code=303)

@app.get("/clients/{client_id}/roster/{roster_id}/download_csv")
async def download_roster_csv(
    client_id: int,
    roster_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    from fastapi.responses import StreamingResponse
    
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    roster = db.query(Roster).filter(Roster.id == roster_id).first()
    if not roster:
        raise HTTPException(status_code=404, detail="Roster not found")
    
    entries = db.query(RosterEntry).filter(RosterEntry.roster_id == roster_id).all()
    
    output = io.StringIO()
    writer = csv.writer(output)
    
    writer.writerow([
        'Primary ID', 'Last Name', 'First Name', 'Company', 'Modality', 
        'Location', 'Division', 'Supervisor Name', 'Alternate Id', 
        'Alternate Id Type', 'Alternate Id 2', 'Alternate Id 2 Type',
        'Employee Name', 'Employee ID', 'Position', 'Department',
        'Testing Status', 'Test Date'
    ])
    
    for entry in entries:
        test_date_str = entry.test_date.strftime('%Y-%m-%d') if entry.test_date else ''
        writer.writerow([
            entry.primary_id or '',
            entry.last_name or '',
            entry.first_name or '',
            entry.company or '',
            entry.modality or '',
            entry.location or '',
            entry.division or '',
            entry.supervisor_name or '',
            entry.alternate_id or '',
            entry.alternate_id_type or '',
            entry.alternate_id_2 or '',
            entry.alternate_id_2_type or '',
            entry.employee_name,
            entry.employee_id or '',
            entry.position or '',
            entry.department or '',
            entry.testing_status,
            test_date_str
        ])
    
    output.seek(0)
    filename = f"{client.name}_roster_{roster.quarter}.csv".replace(' ', '_')
    
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@app.get("/clients/{client_id}/roster/{roster_id}/download_xlsx")
async def download_roster_xlsx(
    client_id: int,
    roster_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    roster = db.query(Roster).filter(Roster.id == roster_id).first()
    if not roster:
        raise HTTPException(status_code=404, detail="Roster not found")
    
    entries = db.query(RosterEntry).filter(RosterEntry.roster_id == roster_id).all()
    
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Roster"
    
    # Add headers with formatting
    headers = [
        'Primary ID', 'Last Name', 'First Name', 'Company', 'Modality', 
        'Location', 'Division', 'Supervisor Name', 'Alternate Id', 
        'Alternate Id Type', 'Alternate Id 2', 'Alternate Id 2 Type',
        'Employee Name', 'Employee ID', 'Position', 'Department',
        'Testing Status', 'Test Date'
    ]
    ws.append(headers)
    
    # Format header row
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    
    # Add data rows
    for entry in entries:
        test_date_str = entry.test_date.strftime('%Y-%m-%d') if entry.test_date else ''
        ws.append([
            entry.primary_id or '',
            entry.last_name or '',
            entry.first_name or '',
            entry.company or '',
            entry.modality or '',
            entry.location or '',
            entry.division or '',
            entry.supervisor_name or '',
            entry.alternate_id or '',
            entry.alternate_id_type or '',
            entry.alternate_id_2 or '',
            entry.alternate_id_2_type or '',
            entry.employee_name,
            entry.employee_id or '',
            entry.position or '',
            entry.department or '',
            entry.testing_status,
            test_date_str
        ])
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Save to bytes
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"{client.name}_roster_{roster.quarter}.xlsx".replace(' ', '_')
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@app.post("/clients/{client_id}/attachments")
async def upload_attachment(
    client_id: int,
    file: UploadFile = File(...),
    category: str = Form("roster_request"),
    description: Optional[str] = Form(None),
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    if not file.filename:
        raise HTTPException(status_code=400, detail="No file provided")
    
    original_filename = os.path.basename(file.filename)
    sanitized_filename = re.sub(r'[^a-zA-Z0-9._-]', '_', original_filename)
    
    content = await file.read()
    file_size = len(content)
    
    if file_size > 3 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File size must be less than 3MB")
    
    unique_id = str(uuid.uuid4())
    extension = os.path.splitext(sanitized_filename)[1]
    stored_filename = f"{client_id}_{unique_id}{extension}"
    file_path = os.path.join("uploads", "attachments", stored_filename)
    
    with open(file_path, "wb") as buffer:
        buffer.write(content)
    
    attachment = Attachment(
        client_id=client_id,
        filename=stored_filename,
        original_filename=original_filename,
        file_size=file_size,
        category=category,
        description=description
    )
    db.add(attachment)
    db.commit()
    
    return RedirectResponse(url=f"/clients/{client_id}", status_code=303)

@app.get("/attachments/{attachment_id}/download")
async def download_attachment(
    attachment_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    from fastapi.responses import FileResponse
    
    attachment = db.query(Attachment).filter(Attachment.id == attachment_id).first()
    if not attachment:
        raise HTTPException(status_code=404, detail="Attachment not found")
    
    file_path = os.path.join("uploads", "attachments", attachment.filename)
    
    abs_file_path = os.path.abspath(file_path)
    abs_upload_dir = os.path.abspath("uploads/attachments")
    
    if not abs_file_path.startswith(abs_upload_dir):
        raise HTTPException(status_code=403, detail="Access denied")
    
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="File not found")
    
    return FileResponse(
        path=file_path,
        filename=attachment.original_filename,
        media_type="application/octet-stream"
    )

@app.post("/attachments/{attachment_id}/delete")
async def delete_attachment(
    attachment_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    attachment = db.query(Attachment).filter(Attachment.id == attachment_id).first()
    if not attachment:
        raise HTTPException(status_code=404, detail="Attachment not found")
    
    client_id = attachment.client_id
    file_path = os.path.join("uploads", "attachments", attachment.filename)
    
    abs_file_path = os.path.abspath(file_path)
    abs_upload_dir = os.path.abspath("uploads/attachments")
    
    if not abs_file_path.startswith(abs_upload_dir):
        raise HTTPException(status_code=403, detail="Access denied")
    
    if os.path.exists(file_path):
        os.remove(file_path)
    
    db.delete(attachment)
    db.commit()
    
    return RedirectResponse(url=f"/clients/{client_id}", status_code=303)

@app.post("/clients/{client_id}/cc_emails")
async def add_cc_email(
    client_id: int,
    email: str = Form(...),
    name: Optional[str] = Form(None),
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    cc_email = CCEmail(
        client_id=client_id,
        email=email.strip(),
        name=name.strip() if name else None
    )
    db.add(cc_email)
    db.commit()
    
    return RedirectResponse(url=f"/clients/{client_id}", status_code=303)

@app.post("/cc_emails/{cc_email_id}/delete")
async def delete_cc_email(
    cc_email_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    cc_email = db.query(CCEmail).filter(CCEmail.id == cc_email_id).first()
    if not cc_email:
        raise HTTPException(status_code=404, detail="CC email not found")
    
    client_id = cc_email.client_id
    db.delete(cc_email)
    db.commit()
    
    return RedirectResponse(url=f"/clients/{client_id}", status_code=303)

@app.get("/templates", response_class=HTMLResponse)
async def list_templates(request: Request, user: User = Depends(require_login), db: Session = Depends(get_db)):
    templates_list = db.query(EmailTemplate).all()
    return templates.TemplateResponse("templates.html", {
        "request": request,
        "user": user,
        "templates": templates_list
    })

@app.get("/templates/new", response_class=HTMLResponse)
async def new_template_form(request: Request, user: User = Depends(require_login)):
    return templates.TemplateResponse("template_form.html", {
        "request": request,
        "user": user,
        "template": None
    })

@app.post("/templates/new")
async def create_template(
    request: Request,
    name: str = Form(...),
    subject_template: str = Form(...),
    body_template: str = Form(...),
    template_type: str = Form(...),
    active: bool = Form(True),
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    template = EmailTemplate(
        name=name,
        subject_template=subject_template,
        body_template=body_template,
        template_type=template_type,
        active=active
    )
    db.add(template)
    db.commit()
    return RedirectResponse(url="/templates", status_code=303)

@app.get("/templates/{template_id}/edit", response_class=HTMLResponse)
async def edit_template_form(
    request: Request,
    template_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    template = db.query(EmailTemplate).filter(EmailTemplate.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")
    
    return templates.TemplateResponse("template_form.html", {
        "request": request,
        "user": user,
        "template": template
    })

@app.post("/templates/{template_id}/delete")
async def delete_template(
    template_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    template = db.query(EmailTemplate).filter(EmailTemplate.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")
    
    db.delete(template)
    db.commit()
    
    return RedirectResponse(url="/templates", status_code=303)

@app.post("/templates/{template_id}/edit")
async def update_template(
    template_id: int,
    name: str = Form(...),
    subject_template: str = Form(...),
    body_template: str = Form(...),
    template_type: str = Form(...),
    active: bool = Form(True),
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    template = db.query(EmailTemplate).filter(EmailTemplate.id == template_id).first()
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")
    
    template.name = name
    template.subject_template = subject_template
    template.body_template = body_template
    template.template_type = template_type
    template.active = active
    db.commit()
    
    return RedirectResponse(url="/templates", status_code=303)

@app.get("/clients/{client_id}/roster", response_class=HTMLResponse)
async def view_roster(
    request: Request,
    client_id: int,
    roster_id: Optional[int] = None,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    client = db.query(Client).filter(Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    if roster_id:
        roster = db.query(Roster).filter(Roster.id == roster_id).first()
    else:
        roster = db.query(Roster).filter(Roster.client_id == client_id).order_by(Roster.received_at.desc()).first()
    
    if not roster:
        return templates.TemplateResponse("roster_view.html", {
            "request": request,
            "user": user,
            "client": client,
            "roster": None,
            "entries": [],
            "stats": {"total": 0, "tested": 0, "not_tested": 0, "excused": 0}
        })
    
    entries = db.query(RosterEntry).filter(RosterEntry.roster_id == roster.id).all()
    
    tested_count = sum(1 for e in entries if e.testing_status == "tested")
    excused_count = sum(1 for e in entries if e.testing_status == "excused")
    not_tested_count = sum(1 for e in entries if e.testing_status == "not_tested")
    stats = {
        "total": len(entries),
        "tested": tested_count,
        "not_tested": not_tested_count,
        "excused": excused_count
    }
    
    return templates.TemplateResponse("roster_view.html", {
        "request": request,
        "user": user,
        "client": client,
        "roster": roster,
        "entries": entries,
        "stats": stats
    })

@app.post("/rosters/{roster_id}/delete")
async def delete_roster(
    roster_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    roster = db.query(Roster).filter(Roster.id == roster_id).first()
    if not roster:
        raise HTTPException(status_code=404, detail="Roster not found")
    
    client_id = roster.client_id
    
    if roster.file_path and os.path.exists(roster.file_path):
        os.remove(roster.file_path)
    
    db.delete(roster)
    db.commit()
    
    return RedirectResponse(url=f"/clients/{client_id}", status_code=303)

@app.post("/roster_entries/{entry_id}/toggle_test")
async def toggle_test_status(
    entry_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    entry = db.query(RosterEntry).filter(RosterEntry.id == entry_id).first()
    if not entry:
        raise HTTPException(status_code=404, detail="Entry not found")
    
    status_cycle = {
        "not_tested": "tested",
        "tested": "excused",
        "excused": "not_tested"
    }
    
    entry.testing_status = status_cycle.get(entry.testing_status, "not_tested")
    entry.test_date = datetime.utcnow() if entry.testing_status == "tested" else None
    db.commit()
    
    return RedirectResponse(url=f"/clients/{entry.client_id}/roster?roster_id={entry.roster_id}", status_code=303)

@app.post("/rosters/{roster_id}/add_entry")
async def add_roster_entry(
    roster_id: int,
    employee_name: str = Form(...),
    employee_id: Optional[str] = Form(None),
    position: Optional[str] = Form(None),
    department: Optional[str] = Form(None),
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    roster = db.query(Roster).filter(Roster.id == roster_id).first()
    if not roster:
        raise HTTPException(status_code=404, detail="Roster not found")
    
    entry = RosterEntry(
        roster_id=roster_id,
        client_id=roster.client_id,
        employee_name=employee_name.strip(),
        employee_id=employee_id.strip() if employee_id else None,
        position=position.strip() if position else None,
        department=department.strip() if department else None,
        testing_status="not_tested"
    )
    db.add(entry)
    db.commit()
    
    return RedirectResponse(url=f"/clients/{roster.client_id}/roster?roster_id={roster_id}", status_code=303)

@app.post("/roster_entries/{entry_id}/delete")
async def delete_roster_entry(
    entry_id: int,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    entry = db.query(RosterEntry).filter(RosterEntry.id == entry_id).first()
    if not entry:
        raise HTTPException(status_code=404, detail="Entry not found")
    
    client_id = entry.client_id
    roster_id = entry.roster_id
    
    db.delete(entry)
    db.commit()
    
    return RedirectResponse(url=f"/clients/{client_id}/roster?roster_id={roster_id}", status_code=303)

@app.get("/logs", response_class=HTMLResponse)
async def email_logs(
    request: Request,
    client_id: Optional[int] = None,
    status: Optional[str] = None,
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    query = db.query(EmailLog).order_by(EmailLog.sent_at.desc())
    
    if client_id:
        query = query.filter(EmailLog.client_id == client_id)
    if status:
        query = query.filter(EmailLog.status == status)
    
    logs = query.limit(100).all()
    clients = db.query(Client).all()
    
    return templates.TemplateResponse("logs.html", {
        "request": request,
        "user": user,
        "logs": logs,
        "clients": clients,
        "selected_client_id": client_id,
        "selected_status": status
    })

@app.get("/settings", response_class=HTMLResponse)
async def settings_page(request: Request, user: User = Depends(require_login), db: Session = Depends(get_db)):
    email_service = EmailService(db)
    config = email_service.get_config()
    
    return templates.TemplateResponse("settings.html", {
        "request": request,
        "user": user,
        "config": config
    })

@app.post("/settings/email")
async def update_email_settings(
    provider: str = Form(...),
    smtp_username: str = Form(...),
    smtp_password: str = Form(...),
    from_email: str = Form(...),
    from_name: str = Form("Random Pool Management"),
    user: User = Depends(require_login),
    db: Session = Depends(get_db)
):
    email_service = EmailService(db)
    
    if provider == "gmail":
        email_service.config.configure_gmail(smtp_username, smtp_password, from_email)
    elif provider == "outlook":
        email_service.config.configure_outlook(smtp_username, smtp_password, from_email)
    
    email_service.config.save_setting("from_name", from_name)
    
    return RedirectResponse(url="/settings", status_code=303)

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=5000)
